from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path("无人机三维地图Web联机躲猫猫项目说明与任务分工.docx")


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def remove_paragraph_borders(container):
    if hasattr(container, "_p"):
        properties = container._p.get_or_add_pPr()
    else:
        properties = container._element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(text)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(text)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.18

for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 11.5)):
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(size)
    style.font.bold = True
    remove_paragraph_borders(style)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(5)
remove_paragraph_borders(title)
title.add_run("无人机三维地图 Web 联机躲猫猫项目")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run("小组项目说明与任务分工")
run.font.name = "Microsoft YaHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(89, 101, 105)

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(10)
intro.add_run("项目结论：").bold = True
intro.add_run("我们小组独立完成一个可以运行和演示的 Web 联机躲猫猫平台。平台使用三维地图作为游戏场景，支持房间创建、多人加入、角色分配、玩家移动、抓捕判定和胜负结算。没有真实航拍地图时，使用内置测试地图或许可证清晰的开源地图，不影响项目完成。")

add_heading(doc, "一 项目做什么", 1)
doc.add_paragraph("我们的作品不是单纯的地图展示页面，也不是无人机飞控软件，而是一个将三维地图转化为多人游戏的 Web 平台。玩家通过浏览器进入房间，在同一张三维地图中进行联机躲猫猫。")
add_bullet(doc, "管理端：选择地图、创建房间、查看房间状态。")
add_bullet(doc, "游戏端：玩家加入房间，在三维场景中移动和寻找。")
add_bullet(doc, "后端：管理房间、同步玩家状态、判断抓捕和结算胜负。")
add_bullet(doc, "地图来源：支持内置测试地图、开源三维地图和后续无人机航拍建模成果。")

add_heading(doc, "二 与无人机课程的关系", 1)
doc.add_paragraph("本项目不要求直接控制无人机，也不把游戏规则建立在飞控功能上。无人机 SDK 或地面站负责航拍任务、影像采集和成果导出；我们的平台负责导入三维成果，配置出生点、躲藏区和边界，并将地图转化为多人联机游戏。")
doc.add_paragraph("答辩时可以这样说明：")
quote = doc.add_paragraph()
quote.paragraph_format.left_indent = Cm(0.6)
quote.paragraph_format.right_indent = Cm(0.4)
quote.paragraph_format.space_after = Pt(8)
quote.add_run("本项目是无人机航拍三维成果的下游游戏化应用。无人机 SDK 负责采集或提供航拍成果，Web 平台负责地图导入、游戏配置和多人联机玩法。")

add_heading(doc, "三 第一版必须完成的功能", 1)
for item in [
    "打开 Web 平台并查看地图列表。",
    "选择内置地图或导入一个 GLB 三维地图。",
    "创建房间并生成房间码，其他玩家可以通过房间码加入。",
    "系统自动分配寻找者和躲藏者。",
    "玩家在三维场景中使用键盘移动。",
    "服务端同步玩家位置和游戏状态。",
    "寻找者接近躲藏者时触发抓捕。",
    "倒计时结束或所有躲藏者被找到后显示结算结果。",
]:
    add_number(doc, item)

add_heading(doc, "四 五人任务分工", 1)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(2.0), Cm(3.2), Cm(7.2), Cm(3.1)]
headers = ["成员", "负责方向", "主要任务", "阶段成果"]
for index, header in enumerate(headers):
    cell = table.rows[0].cells[index]
    cell.width = widths[index]
    set_cell_text(cell, header, bold=True, color="FFFFFF", size=9.5)
    set_cell_shading(cell, "355C63")

rows = [
    ("组长", "系统架构和联机核心", "负责前后端结构、共享数据类型、Express API、Socket.IO、房间状态机、角色分配、位置同步、抓捕判定、胜负结算和最终集成。", "可运行的最终版本和答辩演示"),
    ("队员 A", "三维地图和玩家控制", "负责 GLB 地图加载、场景相机、键盘移动、玩家模型、出生点、躲藏区、地图边界和模型加载异常处理。", "地图加载组件和游戏场景"),
    ("队员 B", "管理端和联机大厅", "负责地图列表、地图导入页面、房间创建、房间码加入、玩家等待页面和前端错误提示。", "管理后台和联机大厅"),
    ("队员 C", "游戏资源和插件", "负责角色、道具、标记、特效和基础资源；整理资源来源和许可证；保证没有外部资源时仍有默认资源可运行。", "默认资源包和资源说明"),
    ("队员 D", "HUD 测试和项目材料", "负责计时器、角色状态、被抓提示、结算页面、多窗口联机测试、异常测试、启动文档和答辩演示材料。", "HUD、测试报告和演示脚本"),
]
for row_index, row in enumerate(rows, start=1):
    cells = table.add_row().cells
    for index, value in enumerate(row):
        cells[index].width = widths[index]
        set_cell_text(cells[index], value, size=9.2)
        if row_index % 2 == 0:
            set_cell_shading(cells[index], "F3F6F5")

add_heading(doc, "五 当前协作方式", 1)
for item in [
    "每个人先确认自己的负责方向，并在群里回复：姓名、负责模块、第一周准备完成的内容。",
    "前端页面通过接口调用后端，不直接修改其他成员负责的模块。",
    "地图文件统一使用 GLB，地图配置单独保存出生点、边界和躲藏区。",
    "外部地图或模型必须记录来源、作者和许可证；许可证不清晰的资源不作为正式成果提交。",
    "每周至少完成一次可运行版本，不能只提交截图或孤立代码。",
]:
    add_number(doc, item)

add_heading(doc, "六 第一阶段安排", 1)
doc.add_paragraph("第一阶段先让项目整体跑起来，暂时不等待真实地图和高级资源。")
first_stage = doc.add_table(rows=1, cols=3)
first_stage.alignment = WD_TABLE_ALIGNMENT.CENTER
first_stage.autofit = False
stage_widths = [Cm(3.2), Cm(7.5), Cm(4.8)]
for index, header in enumerate(["时间", "共同目标", "主要负责人"]):
    first_stage.rows[0].cells[index].width = stage_widths[index]
    set_cell_text(first_stage.rows[0].cells[index], header, bold=True, color="FFFFFF", size=9.5)
    set_cell_shading(first_stage.rows[0].cells[index], "355C63")
stage_rows = [
    ("第 1 周", "确定页面流程、接口格式、地图配置格式和各自开发分支。", "全体，组长汇总"),
    ("第 2 周", "完成前后端骨架、内置测试地图、地图列表和房间 API。", "组长、队员 B"),
    ("第 3 周", "完成联机大厅、房间码加入和三维游戏场景。", "组长、队员 A、队员 B"),
    ("第 4 周", "完成角色分配、移动同步、抓捕判定和基础结算。", "组长、队员 A、队员 D"),
]
for row_index, row in enumerate(stage_rows, start=1):
    cells = first_stage.add_row().cells
    for index, value in enumerate(row):
        cells[index].width = stage_widths[index]
        set_cell_text(cells[index], value, size=9.2)
        if row_index % 2 == 0:
            set_cell_shading(cells[index], "F3F6F5")

add_heading(doc, "七 项目完成标准", 1)
doc.add_paragraph("答辩时能够用两个浏览器窗口完成下面流程，就说明第一版核心目标达成：")
for item in [
    "创建房间并记录房间码。",
    "第二个浏览器输入房间码加入。",
    "开始游戏后看到不同角色和出生点。",
    "移动玩家并看到位置同步。",
    "寻找者靠近躲藏者后触发抓捕。",
    "游戏结束后显示胜负结果。",
]:
    add_bullet(doc, item)

closing = doc.add_paragraph()
closing.paragraph_format.space_before = Pt(10)
closing.paragraph_format.space_after = Pt(0)
closing.add_run("请大家先根据上面的分工回复自己负责的部分。地图和资源可以后续替换，但联机游戏主流程是我们组必须独立完成的核心。")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Drone Hide 项目组  ·  项目说明与任务分工")
footer_run.font.name = "Microsoft YaHei"
footer_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
footer_run.font.size = Pt(8.5)
footer_run.font.color.rgb = RGBColor(130, 140, 140)

doc.core_properties.title = "无人机三维地图 Web 联机躲猫猫项目说明与任务分工"
doc.core_properties.subject = "小组项目说明与任务分工"
doc.core_properties.author = "Drone Hide 项目组"
doc.save(OUTPUT)
print(OUTPUT.resolve())
